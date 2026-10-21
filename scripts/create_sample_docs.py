from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document as DocxDocument
root = Path(__file__).resolve().parents[1] / "samples"
root.mkdir(exist_ok=True)
def make_pdf(path: Path, title: str, lines: list[str]):
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4; y = height - 50
    c.setFont("Helvetica-Bold", 16); c.drawString(50, y, title); y -= 30
    c.setFont("Helvetica", 11)
    for line in lines:
        if y < 60:
            c.showPage(); y = height - 50; c.setFont("Helvetica", 11)
        c.drawString(50, y, line[:110]); y -= 16
    c.save()
def make_docx(path: Path, title: str, lines: list[str]):
    doc = DocxDocument(); doc.add_heading(title, level=1)
    for line in lines: doc.add_paragraph(line)
    doc.save(str(path))
make_pdf(root / "employee_handbook.pdf", "Employee Handbook", ["Probation period is 90 days.", "Employees are eligible for annual leave after completion of the probation period.", "Remote work requires manager approval and attendance during core hours.", "Disciplinary actions follow a progressive review process."])
make_docx(root / "leave_policy.docx", "Leave Policy", ["New joiners receive 12 paid leave days after probation.", "Carry forward is limited to 5 days per calendar year.", "Sick leave requires notification to the manager on the same day.", "Maternity and paternity leave are governed by applicable law."])
make_pdf(root / "recruitment_workflow_guide.pdf", "Recruitment Workflow Guide", ["Each requisition starts with headcount approval.", "The TA team screens candidates within 48 hours of application.", "Interview feedback must be recorded in the ATS within 24 hours.", "Offer approval requires compensation and hiring manager sign-off."])
